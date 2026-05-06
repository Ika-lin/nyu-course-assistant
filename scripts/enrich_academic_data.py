import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
REFERENCE_DIR = ROOT.parent

GLOBAL_COURSES_XLSX = REFERENCE_DIR / "Global Courses Satisfying Shanghai Degree Requirements.xlsx"
CS_EQUIV_XLSX = REFERENCE_DIR / "Computer Science Pre-requisites & Equivalencies.xlsx"


def normalize_code(value):
    if value is None:
        return ""
    text = str(value).replace("\u00ad", "-").replace("\t", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip()
    match = re.search(r"[A-Z]{2,5}-[A-Z]{2,4}\s*\d+[A-Z]?", text)
    if not match:
        return ""
    return re.sub(r"\s+", " ", match.group(0)).strip()


def clean(value):
    if value is None:
        return ""
    text = str(value).replace("\u00ad", "-").replace("\t", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", text).strip()


def first_course_code(value):
    return normalize_code(value)


def extract_global_requirement_courses():
    wb = load_workbook(GLOBAL_COURSES_XLSX, read_only=True, data_only=True)
    records = []
    by_requirement = defaultdict(list)
    by_core = defaultdict(list)
    by_course = defaultdict(list)

    skip_sheets = {"Dropdown Menu", "Instructions ", "Rejected"}
    for ws in wb.worksheets:
        if ws.title in skip_sheets:
            continue

        rows = ws.iter_rows(values_only=True)
        headers = [clean(cell) for cell in next(rows)]
        index = {header: i for i, header in enumerate(headers) if header}

        for row in rows:
            code = normalize_code(row[index.get("Course Number", -1)] if "Course Number" in index else "")
            if not code:
                continue

            title = clean(row[index.get("Course Title", -1)] if "Course Title" in index else "")
            credits_raw = row[index.get("Credit", -1)] if "Credit" in index else None
            try:
                credits = int(float(credits_raw)) if credits_raw not in (None, "") else None
            except (TypeError, ValueError):
                credits = None

            major_minor = clean(row[index.get("Major/Minor", -1)] if "Major/Minor" in index else "")
            requirement = clean(row[index.get("Major Requirement Catagory", -1)] if "Major Requirement Catagory" in index else "")
            core = clean(row[index.get("Core", -1)] if "Core" in index else "")
            gn_minor = clean(row[index.get("GN Minor", -1)] if "GN Minor" in index else "")
            note = clean(row[index.get("Note", -1)] if "Note" in index else "")

            record = {
                "course_code": code,
                "title": title,
                "credits": credits,
                "site": ws.title.strip(),
                "major_minor": major_minor,
                "major_requirement_category": requirement,
                "core": core,
                "gn_minor": gn_minor,
                "note": note,
                "source": GLOBAL_COURSES_XLSX.name,
            }
            records.append(record)
            by_course[code].append(record)
            if requirement:
                by_requirement[requirement].append(code)
            if core:
                by_core[core].append(code)

    def uniq_map(mapping):
        return {key: sorted(set(values)) for key, values in sorted(mapping.items())}

    output = {
        "source": GLOBAL_COURSES_XLSX.name,
        "description": "Approved global courses that can satisfy NYU Shanghai degree, major, minor, core, or GN minor requirements.",
        "records": records,
        "by_requirement": uniq_map(by_requirement),
        "by_core": uniq_map(by_core),
        "by_course": dict(sorted(by_course.items())),
    }
    return output


def extract_cs_equivalencies_prereqs():
    wb = load_workbook(CS_EQUIV_XLSX, read_only=True, data_only=True)
    equivalencies = []
    prereqs = []

    ws = wb["Equivalent Courses Across Schoo"]
    for row in ws.iter_rows(min_row=3, values_only=True):
        cas = clean(row[1] if len(row) > 1 else "")
        tandon = clean(row[2] if len(row) > 2 else "")
        abu_dhabi = clean(row[3] if len(row) > 3 else "")
        shanghai = clean(row[4] if len(row) > 4 else "")
        codes = {
            "cas_courant": first_course_code(cas),
            "tandon": first_course_code(tandon),
            "abu_dhabi": first_course_code(abu_dhabi),
            "shanghai": first_course_code(shanghai),
        }
        if any(codes.values()):
            equivalencies.append({
                "cas_courant": cas,
                "tandon": tandon,
                "abu_dhabi": abu_dhabi,
                "shanghai": shanghai,
                "codes": codes,
                "notes": clean(row[5] if len(row) > 5 else ""),
            })

    sheet_specs = [
        ("Updating CASCourant", 2, "ny_cas_courant"),
        ("Tandon", 3, "ny_tandon"),
        ("Abu Dhabi", 3, "abu_dhabi"),
        ("Shanghai", 3, "shanghai"),
    ]
    for sheet_name, header_row, source in sheet_specs:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=header_row, values_only=True))
        if not rows:
            continue
        headers = [clean(cell).lower() for cell in rows[0]]
        for row in rows[1:]:
            joined = " ".join(clean(cell) for cell in row if clean(cell))
            code = first_course_code(joined)
            if not code:
                continue
            prereqs.append({
                "course_code": code,
                "source_sheet": sheet_name,
                "source_scope": source,
                "course_text": joined,
                "columns": {
                    headers[i] or f"column_{i+1}": clean(value)
                    for i, value in enumerate(row)
                    if i < len(headers) and clean(value)
                },
            })

    return {
        "source": CS_EQUIV_XLSX.name,
        "description": "CS-related cross-campus equivalencies and prerequisite reference rows.",
        "equivalencies": equivalencies,
        "prerequisites": prereqs,
    }


def enrich_courses(requirement_data, cs_data):
    courses_path = DATA_DIR / "courses_complete.json"
    courses = json.loads(courses_path.read_text(encoding="utf-8"))
    title_lookup = {}
    credit_lookup = {}

    for record in requirement_data["records"]:
        code = record["course_code"]
        if record.get("title"):
            title_lookup.setdefault(code, record["title"])
        if record.get("credits"):
            credit_lookup.setdefault(code, record["credits"])

    for item in cs_data["equivalencies"]:
        for text in [item.get("cas_courant"), item.get("tandon"), item.get("abu_dhabi"), item.get("shanghai")]:
            code = first_course_code(text)
            if not code:
                continue
            title = clean(re.sub(r"^[A-Z]{2,5}-[A-Z]{2,4}\s*\d+[A-Z]?\s*", "", text or ""))
            if title:
                title_lookup.setdefault(code, title)

    enriched = 0
    for code, info in courses.items():
        if info.get("title") in ("", "Course", None) and title_lookup.get(code):
            info["title"] = title_lookup[code]
            info["title_source"] = "reference_excel"
            enriched += 1
        if not info.get("credits") and credit_lookup.get(code):
            info["credits"] = credit_lookup[code]

    courses_path.write_text(json.dumps(courses, ensure_ascii=False, indent=2), encoding="utf-8")

    parent_courses_path = REFERENCE_DIR / "courses_complete.json"
    if parent_courses_path.exists():
        parent_courses_path.write_text(json.dumps(courses, ensure_ascii=False, indent=2), encoding="utf-8")

    return enriched


def main():
    requirement_data = extract_global_requirement_courses()
    cs_data = extract_cs_equivalencies_prereqs()

    (DATA_DIR / "study_away_requirement_courses.json").write_text(
        json.dumps(requirement_data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (DATA_DIR / "cs_equivalencies_prereqs.json").write_text(
        json.dumps(cs_data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    enriched = enrich_courses(requirement_data, cs_data)
    print(json.dumps({
        "study_away_records": len(requirement_data["records"]),
        "requirement_categories": len(requirement_data["by_requirement"]),
        "core_categories": len(requirement_data["by_core"]),
        "cs_equivalencies": len(cs_data["equivalencies"]),
        "cs_prerequisite_rows": len(cs_data["prerequisites"]),
        "course_titles_enriched": enriched,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
